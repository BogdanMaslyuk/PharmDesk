"""Imports pharmacy export files; does not invent active ingredients or stock counts."""
import csv
from datetime import datetime
from pathlib import Path
import re
from sqlalchemy import select
from sqlalchemy.orm import Session
from .models import Drawer, Medicine, StockItem, Ingredient

ALIASES = {
    'code': {'код', 'кодтовара', 'артикул', 'articlecode', 'sku'},
    'name': {'наименованиетовара', 'наименование', 'название', 'name'},
    'manufacturer': {'производитель', 'manufacturer', 'фирма'},
    'date': {'датадобавления', 'дата', 'addedon'},
}
# Only demo data: no automatic unverified substance matching for real imports.
DEMO_INGREDIENTS = {
    'Парацетамол учебный': ['Парацетамол'],
    'Ибупрофен учебный': ['Ибупрофен'],
    'Лоратадин учебный': ['Лоратадин'],
    'Цетиризин учебный': ['Цетиризин'],
}


def _key(value):
    return re.sub(r'[\s_\-]+', '', str(value).strip().lower())


def _row_dict(raw):
    names = {_key(k): v for k, v in raw.items() if k is not None}
    return {field: next((names[k] for k in options if k in names), None) for field, options in ALIASES.items()}


def rows_from_file(path):
    path = Path(path)
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = book.active
            iterator = sheet.iter_rows(values_only=True)
            keys = next(iterator)
            for values in iterator:
                if any(value is not None for value in values):
                    yield dict(zip(keys, values))
        finally:
            book.close()
    elif path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as handle:
            sample = handle.read(4096)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=';,\t')
            except csv.Error:
                dialect = csv.excel
            for row in csv.DictReader(handle, dialect=dialect):
                yield row
    else:
        raise ValueError('Поддерживаются только CSV (UTF-8) и XLSX.')


def parse_date(value):
    if not value:
        return None
    if hasattr(value, 'date') and not isinstance(value, str):
        return value.date()
    if hasattr(value, 'year') and not isinstance(value, str):
        return value
    for fmt in ('%d.%m.%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    return None


def import_rows(db: Session, raw_rows, drawer_number: int, drawer_name: str, *, demo=False, commit=True):
    """Upsert by drawer+code; rollback should be handled by caller on failure."""
    if drawer_number <= 0:
        raise ValueError('Номер ящика должен быть положительным.')
    drawer = db.scalar(select(Drawer).where(Drawer.number == drawer_number))
    if drawer is None:
        drawer = Drawer(number=drawer_number, name=drawer_name)
        db.add(drawer)
        db.flush()
    elif drawer_name:
        drawer.name = drawer_name
    created = updated = skipped = 0
    for raw in raw_rows:
        row = _row_dict(raw)
        code_value = row['code']
        code = str(int(code_value)) if isinstance(code_value, float) and code_value.is_integer() else str(code_value or '').strip()
        name = str(row['name'] or '').strip()
        if not code or not name:
            skipped += 1
            continue
        maker = str(row['manufacturer'] or '').strip()
        fingerprint = f'{name.casefold()}::{maker.casefold()}'
        medicine = db.scalar(select(Medicine).where(Medicine.fingerprint == fingerprint))
        if medicine is None:
            medicine = Medicine(fingerprint=fingerprint, name=name, manufacturer=maker)
            db.add(medicine)
            db.flush()
        if demo and not medicine.ingredients:
            for prefix, substance_names in DEMO_INGREDIENTS.items():
                if name.casefold().startswith(prefix.casefold()):
                    for substance_name in substance_names:
                        ingredient = db.scalar(select(Ingredient).where(Ingredient.name == substance_name))
                        if ingredient is None:
                            ingredient = Ingredient(name=substance_name)
                            db.add(ingredient)
                            db.flush()
                        medicine.ingredients.append(ingredient)
                    break
        stock = db.scalar(select(StockItem).where(StockItem.drawer_id == drawer.id, StockItem.article_code == code))
        if stock is None:
            stock = StockItem(article_code=code, drawer=drawer, medicine=medicine, added_on=parse_date(row['date']), source='демо' if demo else 'импорт')
            db.add(stock)
            created += 1
        else:
            stock.medicine = medicine
            stock.added_on = parse_date(row['date'])
            updated += 1
    if commit:
        db.commit()
    else:
        db.flush()
    return {'created': created, 'updated': updated, 'skipped': skipped}
