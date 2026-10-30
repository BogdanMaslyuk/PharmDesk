"""Preview and atomic import of multiple pharmacy drawers from CSV/XLSX.

Never assigns active substances automatically to real pharmacy exports.
"""
import csv
from collections import defaultdict
from io import BytesIO, StringIO
from pathlib import Path
import re

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from .importer import import_rows, parse_date

MAX_BYTES = 4 * 1024 * 1024
MAX_ROWS = 3000
ALIASES = {
    "code": ("код", "код товара", "артикул", "sku", "article_code"),
    "name": ("наименование товара", "наименование", "название", "name"),
    "manufacturer": ("производитель", "manufacturer"),
    "date": ("дата добавления", "дата", "added_on"),
    "drawer": ("ящик", "номер ящика", "ящик №", "drawer", "drawer_number"),
    "drawer_name": ("название ящика", "категория ящика", "drawer_name"),
}


class ImportErrorUser(ValueError):
    """Safe error shown to the user; no server traceback."""


def _normalize(text):
    return re.sub(r"[^\w]+", "", str(text or "").casefold())


def _cell(row, field):
    return next((row[key] for key in row if _normalize(key) in {_normalize(a) for a in ALIASES[field]}), None)


def _read_rows(raw: bytes, filename: str):
    ext = Path(filename).suffix.casefold()
    if ext == ".csv":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ImportErrorUser("CSV-файл должен иметь кодировку UTF-8.") from exc
        header = text.splitlines()[0] if text.splitlines() else ""
        delim = ";" if ";" in header else "\t" if "\t" in header else ","
        return list(csv.DictReader(StringIO(text, newline=""), delimiter=delim))
    if ext == ".xlsx":
        try:
            book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
            try:
                sheet = book.active
                it = sheet.iter_rows(values_only=True)
                headers = next(it, None)
                if not headers:
                    raise ImportErrorUser("В книге Excel нет заголовков.")
                output = []
                for index, values in enumerate(it):
                    if index >= MAX_ROWS:
                        raise ImportErrorUser(f"Максимум {MAX_ROWS} строк за одну загрузку.")
                    if any(value is not None for value in values):
                        output.append(dict(zip(headers, values)))
                return output
            finally:
                book.close()
        except ImportErrorUser:
            raise
        except Exception as exc:
            raise ImportErrorUser("Не удалось прочитать XLSX. Проверьте файл.") from exc
    raise ImportErrorUser("Поддерживаются только файлы .csv (UTF-8) и .xlsx.")


def _int_drawer(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        f = float(str(value).replace(",", "."))
        if not f.is_integer() or f < 1 or f > 999999:
            return None
        return int(f)
    except (TypeError, ValueError):
        return None


def prepare(raw: bytes, filename: str, default_drawer=None, default_name="Без названия"):
    if len(raw) > MAX_BYTES:
        raise ImportErrorUser("Файл превышает 4 МБ.")
    if not raw:
        raise ImportErrorUser("Файл пуст.")
    rows = _read_rows(raw, filename)
    if len(rows) > MAX_ROWS:
        raise ImportErrorUser(f"Максимум {MAX_ROWS} строк за одну загрузку.")
    if not rows:
        raise ImportErrorUser("В файле нет товарных позиций.")
    keys = {_normalize(key) for key in rows[0]}
    for required in ("code", "name"):
        if not any(_normalize(alias) in keys for alias in ALIASES[required]):
            raise ImportErrorUser("В файле обязательны столбцы «Код» и «Наименование товара».")
    groups = defaultdict(list)
    drawer_names = {}
    seen = set()
    errors = []
    valid = 0
    for index, row in enumerate(rows, start=2):
        code_raw = _cell(row, "code")
        code = str(int(code_raw)) if isinstance(code_raw, float) and code_raw.is_integer() else str(code_raw or "").strip()
        name = str(_cell(row, "name") or "").strip()
        manufacturer = str(_cell(row, "manufacturer") or "").strip()
        row_drawer = _cell(row, "drawer")
        drawer = _int_drawer(row_drawer if row_drawer is not None and str(row_drawer).strip() else default_drawer)
        drawer_title = str(_cell(row, "drawer_name") or default_name).strip() or "Без названия"
        date = _cell(row, "date")
        cause = None
        if not code or not name:
            cause = "нет кода или названия"
        elif len(code) > 80 or len(name) > 350 or len(manufacturer) > 300 or len(drawer_title) > 200:
            cause = "слишком длинное значение"
        elif drawer is None:
            cause = "не указан корректный номер ящика"
        elif date is not None and str(date).strip() and parse_date(date) is None:
            cause = "непонятная дата добавления"
        elif (drawer, code) in seen:
            cause = "дубликат кода в том же ящике внутри файла"
        elif drawer in drawer_names and drawer_names[drawer] != drawer_title:
            cause = "разные названия одного ящика в файле"
        if cause:
            errors.append(f"Строка {index}: {cause}")
            continue
        seen.add((drawer, code))
        drawer_names[drawer] = drawer_title
        valid += 1
        groups[drawer].append({
            "Код": code,
            "Наименование товара": name,
            "Производитель": manufacturer,
            "Дата добавления": date,
        })
    return groups, drawer_names, errors, len(rows), valid


def batch_import(db: Session, raw: bytes, filename: str, default_drawer=None,
                 default_name="Без названия", *, preview=True):
    groups, drawer_names, errors, total, valid = prepare(raw, filename, default_drawer, default_name)
    created = updated = skipped = 0
    try:
        for number, rows in groups.items():
            info = import_rows(db, rows, number, drawer_names[number], commit=False)
            created += info["created"]
            updated += info["updated"]
            skipped += info["skipped"]
        if preview:
            db.rollback()
        else:
            db.commit()
    except Exception:
        db.rollback()
        raise
    return {
        "mode": "preview" if preview else "applied",
        "rows": total, "valid": valid, "created": created, "updated": updated,
        "skipped": skipped + len(errors), "drawers": len(groups),
        "errors": errors[:30], "more_errors": max(0, len(errors) - 30),
        "message": "Предпросмотр: изменения не сохранены." if preview else "Импорт завершён.",
    }
